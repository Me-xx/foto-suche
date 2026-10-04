"""Caption-Checker: Verdachtsfälle in der digiKam-DB finden und bereinigen.

Aufruf ohne Parameter (Analyse):
    python -m app.caption_check
        → durchsucht alle Captions (ImageComments, type=1) nach
          Datenmüll und redundanten Inhalten und schreibt eine Excel
          ``descriptError_<Datum-Uhrzeit>.xlsx`` mit den Spalten
          ``merge description`` und ``cleanup description``.

Aufruf mit Excel-Parameter (Anwenden):
    python -m app.caption_check <datei.xlsx> [--dry-run] [--db <pfad>]
        → setzt die markierten Spalten um:
          * ``x`` in ``merge description``: die redundante Caption wird
            per LLM (Qwen2-VL) zu einer Beschreibung zusammengefasst
            und in die digiKam-DB geschrieben.
          * ``x`` in ``cleanup description``: die Caption wird geleert
            (comment = ''), damit der Qwen-Tagger das Bild neu taggt.
        Es wird nur geschrieben, was in der Excel als Filename aufgeführt
        und mit ``x`` markiert ist.  Vor dem ersten Write wird eine
        Sicherungskopie der digikam4.db angelegt.

Der Apply-Modus mit LLM benötigt die Qwen-GPU-Umgebung
(C:\\MeineDateienDesk\\MCP\\mcp-env), der Analyse-Modus läuft in jeder
Umgebung mit openpyxl.
"""
from __future__ import annotations

import argparse
import difflib
import re
import shutil
import sqlite3
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Optional

from .config import load_config
from .digikam import DigiKamReader

# ---------------------------------------------------------------------------
# Heuristiken
# ---------------------------------------------------------------------------

# Software-Artefakte aus EXIF/Datei-Headern — nie sinnvolle Bildbeschreibungen
_ARTIFACT_PATTERNS = [
    r"^created with gimp\b",
    r"^creator:\s*gd-jpeg\b",
    r"^konica minolta\b",
    r"^digital camera\s*$",
    r"^lead technologies\b",
    r"^file written by adobe\b",
    r"^olympus\b.*digital camera",
    r"^casio\b.*qv",
    r"^eastman kodak\b",
]
_ARTIFACT_RE = [re.compile(p, re.IGNORECASE) for p in _ARTIFACT_PATTERNS]

# Mojibake-/Encoding-Müll
_MOJIBAKE_RE = re.compile(r"�|Ã.|â€.|Â\xa0")

# Gebräuchliche deutsche/englische Wörter — eine Caption mit >= 6 Wörtern,
# die KEINES davon enthält, ist mit hoher Wahrscheinlichkeit Datenmüll.
_COMMON_WORDS = set(
    "the a an and or of in on at to with from for is are was were be been "
    "der die das ein eine einer eines dem den des und oder in im auf an mit "
    "von für aus bei ist sind war waren wird werden man es sie er das diese "
    "this that these those image photo picture zeigt zeigt's man woman girl "
    "boy child children people group sky day sun water tree house building "
    "straße street room table front background foreground".split()
)


def _alpha_words(text: str) -> list[str]:
    return re.findall(r"[a-zA-ZäöüÄÖÜß]{3,}", text)


def garbage_reason(caption: str) -> Optional[str]:
    """Liefert einen Grund, wenn die Caption Datenmüll ist, sonst None."""
    if not caption or not caption.strip():
        return "leer/nur Whitespace"
    if any(ord(ch) < 32 and ch not in "\n\r\t" for ch in caption):
        return "Steuerzeichen (z.B. Null-Bytes)"
    if _MOJIBAKE_RE.search(caption):
        return "Mojibake/Encoding-Fehler"
    stripped = caption.strip()
    for rex in _ARTIFACT_RE:
        if rex.search(stripped):
            return "Software-Artefakt (EXIF/Editor)"
    words = _alpha_words(caption)
    if len(stripped) < 40 and len(words) < 5:
        return "bedeutungslos/zu kurz"
    if len(words) >= 6 and not any(w.lower() in _COMMON_WORDS for w in words):
        return "kein erkennbarer deutscher/englischer Text"
    return None


def _split_segments(caption: str) -> list[str]:
    return [s.strip() for s in re.split(r"\s*\|\s*|\r?\n+", caption) if s.strip()]


def _token_jaccard(a: str, b: str) -> float:
    wa = {w.lower() for w in _alpha_words(a)}
    wb = {w.lower() for w in _alpha_words(b)}
    if not wa or not wb:
        return 0.0
    return len(wa & wb) / len(wa | wb)


def redundancy_reason(caption: str) -> Optional[str]:
    """Liefert einen Grund, wenn Segmente der Caption redundant sind."""
    segments = _split_segments(caption)
    if len(segments) < 2:
        return None
    for i in range(len(segments)):
        for j in range(i + 1, len(segments)):
            ratio = difflib.SequenceMatcher(
                None, segments[i].lower(), segments[j].lower()
            ).ratio()
            jac = _token_jaccard(segments[i], segments[j])
            if ratio >= 0.75 or jac >= 0.7:
                return (
                    f"redundante Segmente ({int(max(ratio, jac) * 100)}% "
                    f"Ähnlichkeit: '{segments[j][:40]}…')"
                )
    return None


@dataclass
class Suspect:
    image_id: int
    filename: str
    caption: str
    reason: str


# ---------------------------------------------------------------------------
# Analyse → Excel
# ---------------------------------------------------------------------------

def analyze(cfg: Optional[dict] = None) -> list[Suspect]:
    """Durchsucht alle Captions nach Verdachtsfällen."""
    cfg = cfg or load_config()
    dk = DigiKamReader(
        db_path=cfg["digikam"]["db_path"],
        image_root=cfg["digikam"]["image_root"],
        db_copy_path=cfg["digikam"].get("db_copy_path"),
    )
    suspects: list[Suspect] = []
    try:
        for rec in dk.iter_images(only_with_caption=True):
            reason = garbage_reason(rec.caption or "") or redundancy_reason(
                rec.caption or ""
            )
            if reason:
                suspects.append(
                    Suspect(rec.image_id, rec.path, rec.caption or "", reason)
                )
    finally:
        dk.close()
    suspects.sort(key=lambda s: (s.reason, s.filename.lower()))
    return suspects


# Für openpyxl illegale Zeichen (Steuerzeichen etc.) aus Zellinhalten entfernen
_XLSX_ILLEGAL_RE = re.compile(
    r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff\U0001fffe\U0001ffff]"
)


def _xlsx_safe(text: str) -> str:
    return _XLSX_ILLEGAL_RE.sub("␀", text)


def write_excel(suspects: list[Suspect], out_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Verdachtsfälle"
    headers = [
        "Filename",
        "ImageID",
        "Grund",
        "merge description",
        "cleanup description",
        "Beschreibung",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for s in suspects:
        preview = _xlsx_safe(s.caption)
        if len(preview) > 1000:
            preview = preview[:1000] + "…"
        ws.append([s.filename, s.image_id, s.reason, "", "", preview])

    # Spaltenbreiten für bessere Lesbarkeit
    widths = [50, 10, 42, 18, 18, 90]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    wb.save(out_path)


def run_analyze(cfg: dict) -> Path:
    suspects = analyze(cfg)
    ts = time.strftime("%Y%m%d_%H%M%S")
    out_dir = cfg.get("caption_check", {}).get("excel_dir") or Path(
        __file__
    ).resolve().parent.parent
    out_dir = Path(out_dir)
    out = out_dir / f"descriptError_{ts}.xlsx"
    write_excel(suspects, out)
    print(f"{len(suspects)} Verdachtsfälle gefunden.")
    print(f"Excel geschrieben: {out}")
    print(
        "Zum Anwenden: 'x' in 'merge description' (LLM-Zusammenfassung) oder "
        "'cleanup description' (Caption leeren) eintragen und dann\n"
        "  python -m app.caption_check <excel>  aufrufen."
    )
    return out


# ---------------------------------------------------------------------------
# Excel → digiKam-DB
# ---------------------------------------------------------------------------

@dataclass
class Action:
    image_id: int
    filename: str
    kind: str  # 'merge' | 'cleanup'
    current_caption: Optional[str]


def read_actions(excel_path: Path, db_path: str) -> tuple[list[Action], list[str]]:
    """Liest die markierten Zeilen aus der Excel ein.

    Nur Zeilen mit 'x' in 'merge description' oder 'cleanup description'
    werden als Aktion übernommen.  Zurückgegeben werden zusätzlich
    Warnungen (z.B. Zeilen mit beiden Marken oder nicht gefundenen Bildern).
    """
    import openpyxl

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb.active

    header = {}
    for col, cell in enumerate(ws[1], start=1):
        if cell.value is not None:
            header[str(cell.value).strip()] = col
    for required in ("Filename", "merge description", "cleanup description"):
        if required not in header:
            raise SystemExit(
                f"Spalte '{required}' nicht in der Excel gefunden — "
                f"ist das eine descriptError-Datei?"
            )

    conn = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row

    actions: list[Action] = []
    warnings: list[str] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        filename = row[header["Filename"] - 1]
        if not filename:
            continue
        filename = str(filename).strip()
        merge_x = str(row[header["merge description"] - 1] or "").strip().lower() == "x"
        cleanup_x = str(row[header["cleanup description"] - 1] or "").strip().lower() == "x"
        if not merge_x and not cleanup_x:
            continue
        if merge_x and cleanup_x:
            warnings.append(f"Beide Spalten markiert — übersprungen: {filename}")
            continue

        cur = conn.execute(
            """
            SELECT i.id, ic.comment
              FROM Images i
              JOIN Albums a ON i.album = a.id
             WHERE i.name = ?
               AND (a.relativePath || '/' || ?) LIKE ('%' || ?)
             LIMIT 1
            """,
            (Path(filename).name, Path(filename).name, filename.replace("\\", "/")),
        )
        r = cur.fetchone()
        if r is None:
            # Fallback: exakter Match über ImageID-Spalte, falls vorhanden
            idcol = header.get("ImageID")
            img_id = None
            if idcol and row[idcol - 1]:
                try:
                    img_id = int(row[idcol - 1])
                except (TypeError, ValueError):
                    img_id = None
            if img_id is None:
                warnings.append(f"Bild nicht gefunden — übersprungen: {filename}")
                continue
            r = conn.execute(
                "SELECT i.id, ic.comment FROM Images i "
                "LEFT JOIN ImageComments ic ON ic.imageid = i.id AND ic.type = 1 "
                "WHERE i.id = ?",
                (img_id,),
            ).fetchone()
            if r is None:
                warnings.append(f"ImageID {img_id} nicht gefunden: {filename}")
                continue
        actions.append(
            Action(
                image_id=r["id"],
                filename=filename,
                kind="merge" if merge_x else "cleanup",
                current_caption=r["comment"],
            )
        )
    conn.close()
    return actions, warnings


def _load_llm(model_path: str):
    """Lädt Qwen2-VL (INT4-quantisiert) für Text-Merge — nur im Apply-Modus."""
    import torch
    from transformers import (
        AutoModelForVision2Seq,
        AutoProcessor,
        BitsAndBytesConfig,
    )

    bnb = BitsAndBytesConfig(
        load_in_4bit=True,
        bnb_4bit_use_double_quant=True,
        bnb_4bit_quant_type="nf4",
        bnb_4bit_compute_dtype=torch.float16,
    )
    processor = AutoProcessor.from_pretrained(model_path, trust_remote_code=True)
    model = AutoModelForVision2Seq.from_pretrained(
        model_path,
        quantization_config=bnb,
        device_map="auto",
        low_cpu_mem_usage=True,
        trust_remote_code=True,
    )
    model.eval()
    return processor, model


def merge_caption_llm(processor, model, caption: str) -> str:
    """Fasst eine redundante Caption per LLM zu einer Beschreibung zusammen."""
    import torch

    prompt = (
        "Die folgende Bildbeschreibung enthält mehrere, teils redundante "
        "Einzelbeschreibungen. Fasse sie zu EINER prägnanten, vollständigen "
        "englischen Beschreibung des Bildinhalts zusammen. Entferne "
        "Wiederholungen und nichtssagende Aussagen (z.B. über den Himmel, "
        "wenn kein Himmel zu sehen ist). Gib ausschließlich die neue "
        "Beschreibung als Fließtext aus, ohne Einleitung.\n\n"
        f"Beschreibung:\n{caption}"
    )
    messages = [{"role": "user", "content": prompt}]
    text = processor.apply_chat_template(
        messages, tokenize=False, add_generation_prompt=True
    )
    inputs = processor(text=[text], return_tensors="pt").to(model.device)
    with torch.no_grad():
        out = model.generate(**inputs, max_new_tokens=256, do_sample=False)
    result = processor.batch_decode(
        out[:, inputs["input_ids"].shape[1]:], skip_special_tokens=True
    )[0].strip()
    return result


def _write_caption(conn: sqlite3.Connection, image_id: int, new_comment: str) -> None:
    """Schreibt die Caption (type=1) — wie der Qwen-Tagger."""
    row = conn.execute(
        "SELECT id FROM ImageComments WHERE imageid = ? AND type = 1 LIMIT 1",
        (image_id,),
    ).fetchone()
    if row is not None:
        conn.execute(
            "UPDATE ImageComments SET comment = ? WHERE id = ?",
            (new_comment, row["id"]),
        )
    elif new_comment:
        conn.execute(
            "INSERT INTO ImageComments (imageid, type, language, comment) "
            "VALUES (?, 1, 'x-default', ?)",
            (image_id, new_comment),
        )


def run_apply(
    excel_path: Path,
    cfg: dict,
    dry_run: bool = False,
    db_override: Optional[str] = None,
) -> int:
    db_path = db_override or cfg["digikam"]["db_path"]
    actions, warnings = read_actions(excel_path, db_path)
    for w in warnings:
        print(f"  WARNUNG: {w}")

    n_merge = sum(1 for a in actions if a.kind == "merge")
    n_cleanup = sum(1 for a in actions if a.kind == "cleanup")
    print(f"{len(actions)} Aktionen: {n_merge} Merge, {n_cleanup} Cleanup.")
    if not actions:
        return 0

    processor = model = None
    if n_merge:
        model_path = cfg.get("caption_check", {}).get(
            "llm_model_path", "C:/MeineDateienDesk/MCP/models/Qwen2VL"
        )
        print(f"Lade LLM für Merge: {model_path} (INT4) …")
        processor, model = _load_llm(model_path)

    if dry_run:
        print("--- DRY-RUN: es wird nichts geschrieben ---")

    # Sicherungskopie der DB vor dem ersten Write
    if not dry_run:
        bak = Path(db_path).with_suffix(
            f".{time.strftime('%Y%m%d_%H%M%S')}.bak"
        )
        shutil.copy2(db_path, bak)
        print(f"DB-Backup angelegt: {bak}")

    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    try:
        conn.execute("BEGIN IMMEDIATE")  # bricht ab, wenn digiKam die DB sperrt
    except sqlite3.OperationalError as e:
        conn.close()
        raise SystemExit(
            f"digiKam-Datenbank ist gesperrt ({e}) — digiKam schließen und "
            "erneut versuchen."
        )

    n_done = 0
    try:
        for i, act in enumerate(actions, start=1):
            if act.kind == "cleanup":
                print(f"  [{i}/{len(actions)}] Cleanup: {act.filename}")
                if not dry_run:
                    _write_caption(conn, act.image_id, "")
                n_done += 1
            else:
                if not act.current_caption:
                    print(
                        f"  [{i}/{len(actions)}] Merge übersprungen "
                        f"(keine Caption vorhanden): {act.filename}"
                    )
                    continue
                merged = merge_caption_llm(processor, model, act.current_caption)
                print(f"  [{i}/{len(actions)}] Merge: {act.filename}")
                print(f"      → {merged[:120]}{'…' if len(merged) > 120 else ''}")
                if not dry_run:
                    _write_caption(conn, act.image_id, merged)
                n_done += 1
        if not dry_run:
            conn.commit()
        else:
            conn.rollback()
    finally:
        conn.close()

    print(f"{n_done} Aktionen {'simuliert' if dry_run else 'angewendet'}.")
    print(
        "Hinweis: danach 'python -m app.index --sync-only' und "
        "'--embed-only' laufen lassen, damit Index und Embeddings "
        "aktualisiert werden."
    )
    return 0


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main() -> int:
    parser = argparse.ArgumentParser(
        prog="python -m app.caption_check",
        description=(
            "digiKam-Captions auf Datenmüll/Redundanz prüfen (Excel) und "
            "markierte Bereinigungen anwenden."
        ),
    )
    parser.add_argument(
        "excel",
        nargs="?",
        default=None,
        help="descriptError-Excel; ohne diesen Parameter läuft nur die Analyse.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Nur anzeigen, was passieren würde (keine DB-Änderung).",
    )
    parser.add_argument(
        "--db",
        default=None,
        help="Abweichenden digiKam-DB-Pfad (z.B. Testkopie) verwenden.",
    )
    args = parser.parse_args()

    cfg = load_config()
    if args.excel is None:
        run_analyze(cfg)
    else:
        p = Path(args.excel)
        if not p.exists():
            print(f"Excel nicht gefunden: {p}", file=sys.stderr)
            return 1
        run_apply(p, cfg, dry_run=args.dry_run, db_override=args.db)
    return 0


if __name__ == "__main__":
    sys.exit(main())
